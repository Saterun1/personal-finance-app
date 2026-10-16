"""Тесты бизнес-логики приложения «Учёт личных финансов».

Запуск:  python -m pytest test_finance_classes.py -v
          (или python test_finance_classes.py — без pytest)
"""

from __future__ import annotations

import os
import tempfile
import unittest
from datetime import date

import openpyxl

from finance_classes import (
    ANALYTICS_SHEET,
    HEADER,
    SHEET_NAME,
    Category,
    FinanceManager,
    Transaction,
    TransactionType,
    parse_date,
)


class TempManagerMixin:
    """Создаёт временный Excel-файл для каждого теста."""

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self.path = os.path.join(self._tmpdir.name, "transactions.xlsx")
        self.manager = FinanceManager(self.path)

    def tearDown(self) -> None:
        self._tmpdir.cleanup()

    def add(self, amount, category, date_str, description=""):
        return self.manager.add_transaction(
            Transaction(
                amount=amount,
                category=self.manager.get_category(category),
                date=date_str,
                description=description,
            )
        )


class TestEnumsAndParsing(unittest.TestCase):
    def test_from_text_accepts_both_languages(self):
        self.assertIs(TransactionType.from_text("Доход"), TransactionType.INCOME)
        self.assertIs(TransactionType.from_text("income"), TransactionType.INCOME)
        self.assertIs(TransactionType.from_text("РАСХОД"), TransactionType.EXPENSE)

    def test_from_text_rejects_unknown(self):
        with self.assertRaises(ValueError):
            TransactionType.from_text("Перевод")

    def test_parse_date(self):
        self.assertEqual(parse_date("2025-03-07"), date(2025, 3, 7))
        self.assertEqual(parse_date(date(2024, 12, 31)), date(2024, 12, 31))
        with self.assertRaises(ValueError):
            parse_date("07.03.2025")

    def test_transaction_type_values_are_russian(self):
        self.assertEqual(TransactionType.INCOME.value, "Доход")
        self.assertEqual(TransactionType.EXPENSE.value, "Расход")


class TestCategory(unittest.TestCase):
    def test_accepts_enum_or_string(self):
        self.assertEqual(Category("Продукты", "Расход").type, TransactionType.EXPENSE)
        self.assertEqual(Category("Зарплата", TransactionType.INCOME).type, TransactionType.INCOME)

    def test_strips_and_rejects_empty(self):
        self.assertEqual(Category("  Транспорт  ", TransactionType.EXPENSE).name, "Транспорт")
        with self.assertRaises(ValueError):
            Category("   ", TransactionType.EXPENSE)

    def test_to_row(self):
        transaction = Transaction(1500, Category("Зарплата", TransactionType.INCOME), "2025-01-15")
        self.assertEqual(
            transaction.to_row(), [1500.0, "Зарплата", "Доход", "2025-01-15", ""]
        )


class TestTransaction(unittest.TestCase):
    def setUp(self):
        self.income = Category("Зарплата", TransactionType.INCOME)
        self.expense = Category("Продукты", TransactionType.EXPENSE)

    def test_signed_amount(self):
        self.assertEqual(Transaction(100, self.income).signed_amount, 100)
        self.assertEqual(Transaction(100, self.expense).signed_amount, -100)

    def test_rounds_amount(self):
        self.assertEqual(Transaction(10.567, self.income).amount, 10.57)

    def test_rejects_non_positive_and_non_numeric(self):
        with self.assertRaises(ValueError):
            Transaction(0, self.income)
        with self.assertRaises(ValueError):
            Transaction(-5, self.expense)
        with self.assertRaises(ValueError):
            Transaction("abc", self.income)

    def test_rejects_bad_date_and_category(self):
        with self.assertRaises(ValueError):
            Transaction(10, self.income, "31-12-2024")
        with self.assertRaises(ValueError):
            Transaction(10, "Продукты", "2025-01-01")

    def test_default_date_is_today(self):
        self.assertEqual(Transaction(10, self.income).date, date.today())

    def test_date_str_roundtrip(self):
        transaction = Transaction(10, self.income, "2025-02-03")
        self.assertEqual(transaction.date_str, "2025-02-03")


class TestFinanceManagerBasics(TempManagerMixin, unittest.TestCase):
    def test_default_categories_created(self):
        names = {c.name for c in self.manager.categories}
        self.assertIn("Зарплата", names)
        self.assertIn("Продукты", names)
        self.assertTrue(all(c.type is not None for c in self.manager.categories))

    def test_add_and_delete(self):
        self.add(100, "Зарплата", "2025-01-10", "аванс")
        self.add(50, "Продукты", "2025-01-11", "магазин")
        self.assertEqual(len(self.manager.transactions), 2)

        self.assertTrue(self.manager.delete_transaction(0))
        self.assertEqual(len(self.manager.transactions), 1)
        self.assertEqual(self.manager.transactions[0].description, "магазин")

    def test_delete_out_of_range(self):
        self.assertFalse(self.manager.delete_transaction(5))
        self.assertFalse(self.manager.delete_transaction(-1))

    def test_add_transaction_with_kwargs(self):
        self.manager.add_transaction(amount=250, category="Транспорт", date_="2025-02-02")
        self.assertEqual(len(self.manager.transactions), 1)
        self.assertEqual(self.manager.transactions[0].category.name, "Транспорт")

    def test_add_transaction_unknown_category_raises(self):
        with self.assertRaises(ValueError):
            self.manager.add_transaction(amount=250, category="Криптовалюта")

    def test_add_category_is_idempotent(self):
        first = self.manager.add_category("Хобби", TransactionType.EXPENSE)
        second = self.manager.add_category("Хобби", TransactionType.EXPENSE)
        self.assertIs(first, second)
        self.assertEqual(
            sum(1 for c in self.manager.categories if c.name == "Хобби"), 1
        )

    def test_get_category_case_insensitive(self):
        self.assertIsNotNone(self.manager.get_category("зарплата"))
        self.assertIsNone(self.manager.get_category("Неизвестная"))


class TestAnalytics(TempManagerMixin, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.add(1000, "Зарплата", "2025-01-10", "январь")
        self.add(200, "Продукты", "2025-01-15", "магазин")
        self.add(300, "Продукты", "2025-02-02", "магазин")
        self.add(500, "Зарплата", "2025-02-10", "февраль")

    def test_balance(self):
        self.assertEqual(self.manager.get_balance(), 1000)

    def test_income_and_expense_totals(self):
        self.assertEqual(self.manager.get_income_total(), 1500)
        self.assertEqual(self.manager.get_expense_total(), 500)

    def test_category_summary(self):
        summary = self.manager.get_category_summary()
        self.assertEqual(summary["Зарплата"], 1500)
        self.assertEqual(summary["Продукты"], -500)
        self.assertEqual(sum(summary.values()), 1000)

    def test_monthly_summary(self):
        months = self.manager.get_monthly_summary()
        self.assertEqual(list(months), ["2025-02", "2025-01"])
        self.assertEqual(months["2025-01"]["income"], 1000)
        self.assertEqual(months["2025-01"]["expense"], 200)
        self.assertEqual(months["2025-01"]["balance"], 800)
        self.assertEqual(months["2025-02"]["balance"], 200)

    def test_empty_manager_analytics(self):
        empty = FinanceManager(os.path.join(self._tmpdir.name, "empty.xlsx"))
        self.assertEqual(empty.get_balance(), 0)
        self.assertEqual(empty.get_category_summary(), {})
        self.assertEqual(empty.get_monthly_summary(), {})

    def test_clear(self):
        self.manager.clear()
        self.assertEqual(self.manager.transactions, [])
        self.assertEqual(self.manager.get_balance(), 0)


class TestFiltering(TempManagerMixin, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.add(1000, "Зарплата", "2025-01-10", "оклад")
        self.add(200, "Продукты", "2025-01-15", "продукты на неделю")
        self.add(300, "Транспорт", "2025-02-02", "метро")

    def test_filter_by_type(self):
        incomes = self.manager.get_transactions(type_=TransactionType.INCOME)
        self.assertEqual(len(incomes), 1)
        self.assertEqual(incomes[0].category.name, "Зарплата")

    def test_filter_by_category(self):
        self.assertEqual(len(self.manager.get_transactions(category_name="Продукты")), 1)

    def test_filter_by_month(self):
        self.assertEqual(len(self.manager.get_transactions(month="2025-02")), 1)
        self.assertEqual(len(self.manager.get_transactions(month="2025-01")), 2)
        self.assertEqual(len(self.manager.get_transactions(month="2024-12")), 0)

    def test_search_by_description_and_category(self):
        self.assertEqual(len(self.manager.get_transactions(search="метро")), 1)
        self.assertEqual(len(self.manager.get_transactions(search="продукт")), 1)
        self.assertEqual(len(self.manager.get_transactions(search="зарплата")), 1)
        self.assertEqual(len(self.manager.get_transactions(search="зарплата")), 1)

    def test_combined_filters(self):
        result = self.manager.get_transactions(
            type_=TransactionType.EXPENSE, month="2025-01", search="продукты"
        )
        self.assertEqual(len(result), 1)

    def test_sorted_newest_first(self):
        dates = [t.date for t in self.manager.get_transactions()]
        self.assertEqual(dates, sorted(dates, reverse=True))


class TestPersistence(TempManagerMixin, unittest.TestCase):
    def _write_sheet(self, rows) -> None:
        """Записывает произвольные строки в лист «Операции»."""
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(list(HEADER))
        for row in rows:
            sheet.append(row)
        workbook.save(self.path)

    def _load_sheet(self, name: str = SHEET_NAME):
        workbook = openpyxl.load_workbook(self.path)
        self.addCleanup(workbook.close)
        return workbook[name]

    def test_file_has_expected_sheets_and_rows(self):
        self.add(1000, "Зарплата", "2025-01-10", "оклад")
        self.add(199.99, "Продукты", "2025-01-15", "молоко")

        workbook = openpyxl.load_workbook(self.path)
        self.addCleanup(workbook.close)
        self.assertIn(SHEET_NAME, workbook.sheetnames)
        self.assertIn(ANALYTICS_SHEET, workbook.sheetnames)

        sheet = workbook[SHEET_NAME]
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["Amount", "Category", "Type", "Date", "Description"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[2]],
            [1000, "Зарплата", "Доход", "2025-01-10", "оклад"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[3]],
            [199.99, "Продукты", "Расход", "2025-01-15", "молоко"],
        )

    def test_reload_restores_state(self):
        self.add(1000, "Зарплата", "2025-01-10", "оклад")
        self.add(200, "Продукты", "2025-01-15", "молоко")

        restored = FinanceManager(self.path)
        self.assertEqual(len(restored.transactions), 2)
        self.assertEqual(restored.get_balance(), 800)
        self.assertEqual(restored.transactions[1].date, date(2025, 1, 15))
        self.assertIsInstance(restored.transactions[0].category, Category)

    def test_missing_file_gives_empty_manager(self):
        manager = FinanceManager(os.path.join(self._tmpdir.name, "нет-такого.xlsx"))
        self.assertEqual(manager.transactions, [])
        self.assertEqual(len(manager.categories), 7)

    def test_load_skips_broken_rows(self):
        self._write_sheet(
            [
                [1000.00, "Зарплата", "Доход", "2025-01-10", "оклад"],
                ["не число", "Зарплата", "Доход", "2025-01-10", "битая"],
                [50.00, "Продукты", "Расход", "не дата", "битая дата"],
                [None, "Продукты", "Расход", "2025-01-11", "пустая сумма"],
                [70.00, None, "Расход", "2025-01-12", "нет категории"],
            ]
        )

        manager = FinanceManager(self.path)
        self.assertEqual(len(manager.transactions), 1)
        self.assertEqual(manager.get_balance(), 1000)

    def test_load_creates_unknown_category(self):
        self._write_sheet([[500.00, "Книги", "Расход", "2025-03-01", "Оливер Сакс"]])

        manager = FinanceManager(self.path)
        self.assertIsNotNone(manager.get_category("Книги"))
        self.assertEqual(manager.get_category("Книги").type, TransactionType.EXPENSE)

    def test_load_reads_excel_dates(self):
        from datetime import datetime

        self._write_sheet([[500.00, "Продукты", "Расход", datetime(2025, 3, 1), "дата Excel"]])
        manager = FinanceManager(self.path)
        self.assertEqual(manager.transactions[0].date, date(2025, 3, 1))

    def test_cyrillic_readable_after_save(self):
        self.add(1000, "Зарплата", "2025-01-10", "оклад")
        sheet = self._load_sheet()
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("Зарплата", values)
        self.assertIn("оклад", values)

    def test_save_creates_missing_directory(self):
        nested = os.path.join(self._tmpdir.name, "data", "sub", "transactions.xlsx")
        manager = FinanceManager(nested)
        manager.add_transaction(
            amount=10, category=manager.get_category("Продукты"), date_="2025-01-01"
        )
        self.assertTrue(os.path.exists(nested))

    def test_export_to_other_file(self):
        self.add(1000, "Зарплата", "2025-01-10", "оклад")
        target = os.path.join(self._tmpdir.name, "export.xlsx")
        self.manager.save_to_file(target)
        self.assertTrue(os.path.exists(target))
        self.assertEqual(len(FinanceManager(target).transactions), 1)

    def test_analytics_sheet_contains_summary(self):
        self.add(1000, "Зарплата", "2025-01-10", "оклад")
        self.add(200, "Продукты", "2025-01-15", "молоко")

        sheet = self._load_sheet(ANALYTICS_SHEET)
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("Зарплата", values)
        self.assertIn("Продукты", values)
        self.assertIn("Динамика по месяцам", values)
        self.assertIn(800, values)  # итоговый баланс

    def test_corrupted_file_gives_empty_manager(self):
        with open(self.path, "w", encoding="utf-8") as file:
            file.write("это не Excel-файл")
        manager = FinanceManager(self.path)
        self.assertEqual(manager.transactions, [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
